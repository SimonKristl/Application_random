import random
import openpyxl as xl
wb = xl.load_workbook("improvers.xlsx")
sheet = wb.active
# sheet = wb.get_sheet_by_name("sheet")
# sheet = wb.worksheets[0]
from tkinter import *
# setings Tkinter
window = Tk()
window.minsize(800,700)
window.title("Losovaci aplikace pro 3i")
window.resizable(False,False)

# window
img = PhotoImage(file ='light.png')
main_font = ("helvetica", 12)
color1 = "#ffe166"
window.config(bg=color1)
Label(window, image=img, anchor=NW, bg= color1).pack()
# work lists
x = sheet.max_row
list2_all = []
list1_upload = []
list3_random =[]
list4_random = []
control_number = []
winner_number = []
list_control_double_winner = []
list3=[]
r = 0
# read xlsx, extract rows, insert list_box
def upload():
    y = 0
    for _ in range(0, x):
        for cell in list(sheet.rows)[y]:
            list1_upload.append(cell.value)
            list2_all.append(cell.value)
        y += 1
        print(list1_upload)
        list_to = " ".join([str(elam) for elam in list1_upload])
        player_list.insert(END, list_to)
        list1_upload.clear()
        print(list2_all)
        label2.config(text = "Zlepšovatelů celkem: " + str(x) )
# random row from xlsx, control winners, insert to winner_box, control double winner
def random_winner ():
    try:
        r = random.randint(0, x - 1)
        if r in control_number:
            print("Výherce už vyhrál(kontrola řádku)")
            random_winner()

        else:

            control_number.append(r)
            for _ in range(0, 1):
                for cell in list(sheet.rows)[r]:
                    list3_random.append(cell.value)
                    list_to2 = " ".join([str(elam) for elam in list3_random])
                # winer_box.insert(END, list_to2)
                list4_random.append(list(list3_random))
                control_double()
                print(list3)
                for _ in list3:
                    print(_)
                    if _ == 1:
                        print ("funkce if bool")
                        list3.clear()
                        print(list3)
                        list4_random.clear()
                        list3_random.clear()
                        random_winner()
                    else:
                        winer_box.insert(END, list_to2)
                        list_control_double_winner.append(list(list3_random))
                        winner_number.append(list_to2)
                        list3_random.clear()
                        list4_random.clear()
                        w = len(winner_number)
                        label1.config(text="Výherců celkem: " + str(w))
                        list3.clear()
    except:
        print("Chyba v náhodné funkci!")

# delete all
def delete_list ():
    winer_box.delete(0, END)
    player_list.delete(0, END)
    control_number.clear()
    list1_upload.clear()
    list2_all.clear()
    list3_random.clear()
    winner_number.clear()
    list_control_double_winner.clear()
    label1.config(text="Výherci")
    label2.config(text="Zlepšovatelé")
# save result to txt
def save_winner():
    with open ("Vítězové.txt", "w") as file:
        x1=1
        my_task = winer_box.get(0, END)
        for _ in my_task:
            file.write(f'{x1}. {_} ''\n')
            x1 += 1
# check if person didn't win 2x
def control_double ():
    for list in list4_random:
        if list in list_control_double_winner:
            print("Zlepšovatel už jednou vyhrál")
            list3.append(1)
        else:
            print("Zlepšovatel byl vylosován poprvé")
            list3.append(0)
# window frame
text_frame = Frame(window, bg= color1,)
button_frame = Frame(window, bg = color1)
button_frame2 = Frame(window, bg= color1)
text_frame.pack()
button_frame.pack()
button_frame2.pack()

# text
label1 = Label(text_frame, text= "Výherci", font=(main_font, 15), bg= color1)
label1.grid(row=0, column=0)
label2 = Label(text_frame, text="Zlepšovatelé", font=(main_font, 15),bg= color1)
label2.grid(row=0, column=1)
text_scrollbar = Scrollbar(text_frame)
text_scrollbar.grid(row = 1, column = 2, sticky= N+S)
winer_box = Listbox(text_frame, height=20, width=35, borderwidth= 2, font=main_font, )
winer_box.grid(row=1, column=0)
player_list =Listbox(text_frame, height=20, width=35, borderwidth=2, font=main_font, yscrollcommand=text_scrollbar.set)
player_list.grid(row=1, column=1)
text_scrollbar.config(command=player_list.yview)

# buttons
lottery_button = Button(button_frame,text="Vylosovat", borderwidth=4,font= main_font, command=random_winner)
clear_butoon = Button(button_frame2, text="Vymazat seznam", borderwidth=2, font=main_font, command=delete_list)
save_button = Button(button_frame2, text="Uložit výherce", borderwidth=2, font=main_font, command=save_winner)
close_button = Button(button_frame2, text="Zavřít", borderwidth=2, font=main_font, command=window.destroy)
improver_button = Button(button_frame, text="Načti zlepšovatelé", borderwidth=4, font=main_font, command=upload)
lottery_button.grid(row=0,column=1, pady = 10, padx = 20)
clear_butoon.grid(row=2, column=1, padx = 5)
save_button.grid(row=2, column=2, padx=5)
close_button.grid(row=2, column=3, padx = 5)
improver_button.grid(row=0, column=3, pady= 10, padx = 20)

window.mainloop()
